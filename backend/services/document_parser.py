"""
Document Parser Service — Pre-processing layer trước khi gửi vào LightRAG.

Thay vì gửi file_path sang LightRAG để nó tự parse (bị mất text),
backend tự parse trước và gửi text thuần qua POST /documents/text.

Fallback chain:
  DOCX → python-docx DOM traversal → (fallback) LightRAG native
  PDF  → pdfplumber (bảng) + pymupdf (text) → (fallback) LightRAG native
  XLSX → openpyxl → (fallback) LightRAG native
  PPTX → LightRAG native (chưa implement custom parser)
  Other → LightRAG native (metadata-only extension handled upstream)
"""

from __future__ import annotations

import asyncio
from dataclasses import dataclass, field
from pathlib import Path

from backend.logger import logger

# Supported extensions for pre-processing (others fall back to LightRAG native)
PREPROCESS_EXTENSIONS = {".docx", ".doc", ".pdf", ".xlsx", ".xls"}

# Extensions that should only store metadata, never parse content
METADATA_ONLY_EXTENSIONS = {".dwg", ".dxf", ".png", ".jpg", ".jpeg", ".mp4", ".avi", ".step", ".stl"}


@dataclass
class ParseResult:
    text: str
    parser_used: str
    char_count: int = 0
    table_count: int = 0
    success: bool = True
    error: str | None = None

    def __post_init__(self) -> None:
        self.char_count = len(self.text)


def _extract_table_as_markdown(table) -> str:
    """Convert python-docx table to Markdown format."""
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        rows.append("| " + " | ".join(cells) + " |")
    if not rows:
        return ""
    sep = "| " + " | ".join(["---"] * len(table.columns)) + " |"
    rows.insert(1, sep)
    return "\n".join(rows)


def _parse_docx(file_path: Path) -> ParseResult:
    """
    python-docx DOM traversal — đọc theo thứ tự XML, không bỏ sót text giữa ảnh.
    Tốt hơn LightRAG native parser ~2x về completeness.
    """
    try:
        import docx  # python-docx
    except ImportError:
        return ParseResult(text="", parser_used="docx_failed", success=False,
                           error="python-docx not installed")

    try:
        doc = docx.Document(file_path)
        blocks: list[str] = []
        table_index = 0
        para_index = 0
        table_count = 0

        for child in doc.element.body:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

            if tag == "p":
                if para_index < len(doc.paragraphs):
                    para = doc.paragraphs[para_index]
                    para_index += 1
                    text = para.text.strip()
                    if text:
                        blocks.append(text)

            elif tag == "tbl":
                if table_index < len(doc.tables):
                    table = doc.tables[table_index]
                    table_index += 1
                    md = _extract_table_as_markdown(table)
                    if md:
                        table_count += 1
                        blocks.append(f"\n[Bảng {table_count}]\n{md}")

        plain_text = "\n\n".join(blocks)
        return ParseResult(
            text=plain_text,
            parser_used="python-docx",
            table_count=table_count,
        )
    except Exception as exc:
        return ParseResult(text="", parser_used="python-docx", success=False,
                           error=str(exc))


def _parse_pdf(file_path: Path) -> ParseResult:
    """
    pdfplumber (bảng) + pymupdf (text) — combo tốt nhất cho PDF text-based.
    pdfplumber detect cell boundaries, pymupdf preserve paragraph structure.
    """
    try:
        import pdfplumber
        import fitz  # pymupdf
    except ImportError:
        return ParseResult(text="", parser_used="pdf_failed", success=False,
                           error="pdfplumber or pymupdf not installed")

    try:
        blocks_by_page: dict[int, list[str]] = {}
        table_count = 0

        # Pass 1: pdfplumber — extract bảng theo page
        with pdfplumber.open(file_path) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                tables = page.extract_tables() or []
                for table in tables:
                    if not table:
                        continue
                    table_count += 1
                    rows = []
                    for row in table:
                        cells = [str(c or "").strip().replace("\n", " ") for c in row]
                        rows.append("| " + " | ".join(cells) + " |")
                    if rows:
                        sep = "| " + " | ".join(["---"] * len(table[0])) + " |"
                        rows.insert(1, sep)
                        if page_num not in blocks_by_page:
                            blocks_by_page[page_num] = []
                        blocks_by_page[page_num].append(
                            f"\n[Bảng trang {page_num}]\n" + "\n".join(rows)
                        )

        # Pass 2: pymupdf — extract text theo block vật lý
        doc = fitz.open(file_path)
        for page_num in range(doc.page_count):
            page = doc[page_num]
            text = page.get_text("text", sort=True)
            if text and text.strip():
                # Skip mục lục (nhiều dòng kết thúc bằng số trang)
                lines = text.strip().split("\n")
                non_toc = [
                    l for l in lines
                    if not (l.strip() and l.strip()[-1].isdigit()
                            and len(l.strip()) < 60 and "." * 3 in l)
                ]
                clean = "\n".join(non_toc).strip()
                if clean:
                    pg = page_num + 1
                    if pg not in blocks_by_page:
                        blocks_by_page[pg] = []
                    blocks_by_page[pg].insert(0, clean)  # text trước bảng
        doc.close()

        # Ghép theo thứ tự trang
        all_blocks: list[str] = []
        for pg in sorted(blocks_by_page.keys()):
            all_blocks.extend(blocks_by_page[pg])

        plain_text = "\n\n".join(all_blocks)
        return ParseResult(
            text=plain_text,
            parser_used="pdfplumber+pymupdf",
            table_count=table_count,
        )
    except Exception as exc:
        return ParseResult(text="", parser_used="pdfplumber+pymupdf", success=False,
                           error=str(exc))


def _parse_xlsx(file_path: Path) -> ParseResult:
    """openpyxl — extract tất cả worksheets thành Markdown tables."""
    try:
        import openpyxl
    except ImportError:
        return ParseResult(text="", parser_used="xlsx_failed", success=False,
                           error="openpyxl not installed")

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        blocks: list[str] = []
        table_count = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            # Filter empty rows
            non_empty = [r for r in rows if any(c is not None for c in r)]
            if not non_empty:
                continue

            table_count += 1
            md_rows = []
            for i, row in enumerate(non_empty):
                cells = [str(c).strip() if c is not None else "" for c in row]
                md_rows.append("| " + " | ".join(cells) + " |")
                if i == 0:
                    sep = "| " + " | ".join(["---"] * len(row)) + " |"
                    md_rows.append(sep)

            blocks.append(f"## Sheet: {sheet_name}\n\n" + "\n".join(md_rows))

        wb.close()
        plain_text = "\n\n".join(blocks)
        return ParseResult(
            text=plain_text,
            parser_used="openpyxl",
            table_count=table_count,
        )
    except Exception as exc:
        return ParseResult(text="", parser_used="openpyxl", success=False,
                           error=str(exc))


def parse_document(file_path: str | Path, correlation_id: str = "") -> ParseResult | None:
    """
    Main entry point. Returns ParseResult nếu file được pre-process,
    None nếu nên fallback về LightRAG native (extension không support).

    Caller nên check result.success — nếu False thì fallback về native.
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext in METADATA_ONLY_EXTENSIONS:
        # Không parse content — caller xử lý metadata-only
        return None

    if ext not in PREPROCESS_EXTENSIONS:
        # Fallback về LightRAG native
        return None

    logger.info(
        "parser_started",
        file=path.name,
        extension=ext,
        correlation_id=correlation_id,
    )

    if ext in {".docx", ".doc"}:
        result = _parse_docx(path)
    elif ext == ".pdf":
        result = _parse_pdf(path)
    elif ext in {".xlsx", ".xls"}:
        result = _parse_xlsx(path)
    else:
        return None

    if result.success:
        logger.info(
            "parser_completed",
            file=path.name,
            parser=result.parser_used,
            chars=result.char_count,
            tables=result.table_count,
            correlation_id=correlation_id,
        )
    else:
        logger.warning(
            "parser_failed",
            file=path.name,
            parser=result.parser_used,
            error=result.error,
            correlation_id=correlation_id,
        )

    return result


async def parse_document_async(file_path: str | Path, correlation_id: str = "") -> ParseResult | None:
    """Async wrapper — chạy parse trong thread pool để không block event loop."""
    loop = asyncio.get_event_loop()
    return await loop.run_in_executor(None, parse_document, file_path, correlation_id)
