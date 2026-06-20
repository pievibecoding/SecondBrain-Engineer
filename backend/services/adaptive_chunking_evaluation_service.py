from __future__ import annotations

import json
import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from backend.config import settings
from backend.schemas.evaluations import AdaptiveChunkingRunRequest
from backend.services.evaluation_report_service import REPORT_ID_RE, _artifact_dir, get_adaptive_chunking_report


def _safe_run_id(value: str | None) -> str:
    candidate = (value or f"adaptive-chunking-{datetime.now(UTC).strftime('%Y%m%d-%H%M%S')}").strip()
    candidate = re.sub(r"[^A-Za-z0-9_.-]+", "-", candidate).strip("-")
    if not candidate:
        candidate = f"adaptive-chunking-{datetime.now(UTC).strftime('%Y%m%d-%H%M%S')}"
    if not REPORT_ID_RE.match(candidate):
        raise ValueError("Invalid run_id")
    return candidate


def _workspace_root() -> Path:
    return Path(__file__).resolve().parents[2]


def _resolve_source_path(source_path: str) -> Path:
    raw = source_path.strip()
    if not raw:
        raise FileNotFoundError("Source path is empty")

    normalized = raw.replace("\\", "/")
    candidates: list[Path] = []

    if "/local-nas/" in normalized:
        suffix = normalized.split("/local-nas/", 1)[1]
        candidates.append(Path("/local-nas") / suffix)
        candidates.append(_workspace_root() / "local-nas" / suffix)
    elif normalized.startswith("local-nas/"):
        suffix = normalized.removeprefix("local-nas/")
        candidates.append(Path("/local-nas") / suffix)
        candidates.append(_workspace_root() / "local-nas" / suffix)
    elif normalized.startswith("/"):
        candidates.append(Path(normalized))
    else:
        candidates.append(_workspace_root() / raw)

    candidates.append(Path(raw))

    for candidate in candidates:
        try:
            resolved = candidate.resolve()
        except OSError:
            continue
        if resolved.exists():
            return resolved

    raise FileNotFoundError(f"Source path not found: {source_path}")


def _read_pdf(path: Path) -> str:
    try:
        import pdfplumber  # type: ignore
    except Exception as exc:  # pragma: no cover - optional runtime dependency
        raise RuntimeError("pdfplumber is not installed in backend") from exc

    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text(x_tolerance=1, y_tolerance=3) or "")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    try:
        from docx import Document  # type: ignore
    except Exception as exc:  # pragma: no cover - optional runtime dependency
        raise RuntimeError("python-docx is not installed in backend") from exc

    document = Document(path)
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"\n[Table {table_index}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _load_source_text(source_path: str) -> dict[str, Any]:
    path = _resolve_source_path(source_path)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        parser = "text"
    elif suffix == ".pdf":
        text = _read_pdf(path)
        parser = "pdfplumber"
    elif suffix == ".docx":
        text = _read_docx(path)
        parser = "python-docx"
    else:
        raise ValueError(f"Unsupported source extension: {suffix}")

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = text.splitlines()
    empty_lines = sum(1 for line in lines if not line.strip())
    broken_layout_terms = ["cirtceleotohP", "srosneS", "Bảng trang", "| --- |"]
    return {
        "source_path": str(path),
        "parser": parser,
        "text": text,
        "stats": {
            "char_count": len(text),
            "line_count": len(lines),
            "empty_line_ratio": round(empty_lines / max(len(lines), 1), 6),
            "broken_layout_hits": sum(text.count(term) for term in broken_layout_terms),
        },
    }


def _chunk(strategy: str, index: int, text: str, source_path: str, section_hint: str | None = None) -> dict[str, Any]:
    stripped = text.strip()
    return {
        "strategy": strategy,
        "chunk_index": index,
        "text": stripped,
        "char_count": len(stripped),
        "metadata": {"source_path": source_path, "section_hint": section_hint},
    }


def _fixed_chars(text: str, source_path: str, target_chars: int, overlap_chars: int) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    step = max(target_chars - overlap_chars, 1)
    for index, start in enumerate(range(0, len(text), step)):
        part = text[start : start + target_chars]
        if part.strip():
            chunks.append(_chunk("fixed_chars", index, part, source_path))
    return chunks


def _paragraph_merge(text: str, source_path: str, target_chars: int, max_chars: int) -> list[dict[str, Any]]:
    paragraphs = [item.strip() for item in re.split(r"\n\s*\n", text) if item.strip()]
    chunks: list[dict[str, Any]] = []
    current: list[str] = []
    current_len = 0

    for paragraph in paragraphs:
        next_len = current_len + len(paragraph) + 2
        if current and next_len > target_chars:
            chunks.append(_chunk("paragraph_merge", len(chunks), "\n\n".join(current), source_path))
            current = []
            current_len = 0
        if len(paragraph) > max_chars:
            for fixed in _fixed_chars(paragraph, source_path, target_chars=target_chars, overlap_chars=0):
                fixed["strategy"] = "paragraph_merge"
                fixed["chunk_index"] = len(chunks)
                chunks.append(fixed)
            continue
        current.append(paragraph)
        current_len += len(paragraph) + 2

    if current:
        chunks.append(_chunk("paragraph_merge", len(chunks), "\n\n".join(current), source_path))
    return chunks


def _heading_table_aware(text: str, source_path: str, target_chars: int, max_chars: int) -> list[dict[str, Any]]:
    lines = text.splitlines()
    blocks: list[str] = []
    current: list[str] = []

    def flush() -> None:
        nonlocal current
        if current:
            blocks.append("\n".join(current).strip())
            current = []

    for line in lines:
        stripped = line.strip()
        is_heading = bool(stripped) and len(stripped) < 90 and (
            stripped.isupper()
            or re.match(r"^(\d+\.|[A-Z][A-Za-z0-9 -]+ Series\b)", stripped)
            or stripped.endswith("Series")
        )
        is_table = "|" in stripped or stripped.startswith("[Bảng") or stripped.startswith("[Table")
        if is_heading and current and len("\n".join(current)) > target_chars * 0.35:
            flush()
        current.append(line)
        if not is_table and len("\n".join(current)) >= target_chars:
            flush()
        if len("\n".join(current)) >= max_chars:
            flush()
    flush()

    chunks: list[dict[str, Any]] = []
    for block in blocks:
        if len(block) <= max_chars:
            hint = next((line.strip() for line in block.splitlines() if line.strip()), None)
            chunks.append(_chunk("heading_table_aware", len(chunks), block, source_path, hint))
        else:
            for item in _paragraph_merge(block, source_path, target_chars=target_chars, max_chars=max_chars):
                item["strategy"] = "heading_table_aware"
                item["chunk_index"] = len(chunks)
                chunks.append(item)
    return chunks


def _count_hits(text: str, terms: list[str]) -> int:
    lower = text.lower()
    return sum(lower.count(term.lower()) for term in terms if term)


def _contains_any(text: str, terms: list[str]) -> bool:
    lower = text.lower()
    return any(term.lower() in lower for term in terms if term)


def _top_evidence(chunks: list[dict[str, Any]], terms: list[str], limit: int = 5) -> list[dict[str, Any]]:
    scored: list[tuple[int, dict[str, Any]]] = []
    for chunk in chunks:
        text = str(chunk.get("text", ""))
        score = _count_hits(text, terms)
        if score > 0:
            scored.append((score, chunk))
    scored.sort(key=lambda item: (-item[0], int(item[1].get("char_count", 0))))
    return [
        {
            "score": score,
            "strategy": chunk.get("strategy"),
            "chunk_index": chunk.get("chunk_index"),
            "char_count": chunk.get("char_count"),
            "text": str(chunk.get("text", ""))[:2500],
        }
        for score, chunk in scored[:limit]
    ]


def _evaluate_chunks(
    chunks: list[dict[str, Any]],
    candidate_terms: dict[str, list[str]],
    expected_terms: list[str],
    noise_terms: list[str],
) -> dict[str, Any]:
    joined = "\n\n".join(str(chunk.get("text", "")) for chunk in chunks)
    sizes = [len(str(chunk.get("text", ""))) for chunk in chunks]
    by_candidate = {
        candidate: _contains_any(joined, aliases)
        for candidate, aliases in candidate_terms.items()
    }
    candidate_total = max(len(by_candidate), 1)
    by_expected = {term: term.lower() in joined.lower() for term in expected_terms if term}
    expected_total = max(len(by_expected), 1)
    table_lines = sum(1 for line in joined.splitlines() if "|" in line)
    broken_table_lines = sum(1 for line in joined.splitlines() if line.strip() in {"| |", "| --- |"} or "cirtceleotohP" in line)
    candidate_terms_flat = [term for aliases in candidate_terms.values() for term in aliases]

    return {
        "chunk_count": len(chunks),
        "avg_chunk_chars": round(sum(sizes) / max(len(sizes), 1), 2),
        "max_chunk_chars": max(sizes) if sizes else 0,
        "candidate_coverage": {
            "rate": round(sum(1 for value in by_candidate.values() if value) / candidate_total, 6),
            "covered": sum(1 for value in by_candidate.values() if value),
            "total": len(by_candidate),
            "by_candidate": by_candidate,
        },
        "expected_term_coverage": {
            "rate": round(sum(1 for value in by_expected.values() if value) / expected_total, 6),
            "covered": sum(1 for value in by_expected.values() if value),
            "total": len(by_expected),
            "by_term": by_expected,
        },
        "noise_hit_count": _count_hits(joined, noise_terms),
        "table_fragment_ratio": round(broken_table_lines / max(table_lines, 1), 6),
        "top_evidence": _top_evidence(chunks, candidate_terms_flat + expected_terms),
        "top_noise": _top_evidence(chunks, noise_terms, limit=3),
    }


def _score_metrics(metrics: dict[str, Any], max_chars: int) -> dict[str, Any]:
    candidate_rate = float(metrics["candidate_coverage"]["rate"])
    expected_rate = float(metrics["expected_term_coverage"]["rate"])
    noise = int(metrics["noise_hit_count"])
    table_ratio = float(metrics["table_fragment_ratio"])
    max_chunk_chars = int(metrics["max_chunk_chars"])
    oversize_penalty = 25 if max_chunk_chars > max_chars else 0
    normalized_noise = min(noise / 50, 1)
    score = candidate_rate * 40 + expected_rate * 25 - normalized_noise * 20 - table_ratio * 10 - oversize_penalty
    return {
        "score": round(score, 6),
        "rationale": {
            "candidate_rate": candidate_rate,
            "expected_rate": expected_rate,
            "normalized_noise": round(normalized_noise, 6),
            "table_fragment_ratio": table_ratio,
            "oversize_penalty": oversize_penalty,
        },
    }


def _write_json_report(report: dict[str, Any], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{report['run_id']}.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_markdown_report(report: dict[str, Any], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"{report['run_id']}.md"
    lines = [
        f"# {report.get('title', 'Adaptive Chunking Evaluation')}",
        "",
        f"**Run ID:** `{report['run_id']}`  ",
        f"**Recommendation:** `{report['summary']['recommendation']}`",
        "",
        "## Aggregate",
        "",
        "| Metric | Value |",
        "|--------|-------|",
        f"| Documents | {report['summary']['document_count']} |",
        f"| Evaluable documents | {report['summary']['evaluable_document_count']} |",
        f"| Average candidate coverage delta | {report['summary']['avg_candidate_coverage_delta']} |",
        f"| Average noise delta | {report['summary']['avg_noise_delta']} |",
        "",
        "## Documents",
        "",
    ]
    for document in report["documents"]:
        lines.extend([
            f"### {document['document_id']}",
            "",
            f"**Source:** `{document.get('source_path', '-')}`  ",
            f"**Winner:** `{document.get('winner', '-')}`  ",
            f"**Status:** `{document.get('status', '-')}`",
            "",
        ])
        if document.get("error"):
            lines.extend([f"**Error:** {document['error']}", ""])
            continue
        lines.extend([
            "| Strategy | Score | Candidate Coverage | Expected Coverage | Noise Hits | Max Chunk Chars |",
            "|----------|-------|--------------------|-------------------|------------|-----------------|",
        ])
        for strategy in document.get("strategies", []):
            metrics = strategy["metrics"]
            lines.append(
                f"| {strategy['name']} | {strategy['score']['score']} | "
                f"{metrics['candidate_coverage']['rate']} | "
                f"{metrics['expected_term_coverage']['rate']} | "
                f"{metrics['noise_hit_count']} | {metrics['max_chunk_chars']} |"
            )
        lines.extend(["", "#### Top Evidence", ""])
        winner = next((item for item in document.get("strategies", []) if item["name"] == document.get("winner")), None)
        evidence = (winner or {}).get("metrics", {}).get("top_evidence", [])
        if not evidence:
            lines.append("_No evidence chunks found._")
        for item in evidence[:3]:
            lines.extend([
                f"- score `{item['score']}`, chunk `{item['chunk_index']}`, chars `{item['char_count']}`",
                "",
                "```text",
                item["text"][:1000],
                "```",
                "",
            ])
    path.write_text("\n".join(lines), encoding="utf-8")


def run_adaptive_chunking_evaluation(payload: AdaptiveChunkingRunRequest) -> dict[str, Any]:
    run_id = _safe_run_id(payload.run_id)
    output_dir = _artifact_dir()
    config = {
        "target_chars": payload.target_chars,
        "overlap_chars": payload.overlap_chars,
        "max_chars": payload.max_chars,
    }
    documents: list[dict[str, Any]] = []

    for index, document in enumerate(payload.documents, start=1):
        document_id = document.document_id or Path(document.source_path).stem or f"document-{index}"
        try:
            loaded = _load_source_text(document.source_path)
            source_path = str(loaded["source_path"])
            strategies = {
                "fixed_chars": _fixed_chars(loaded["text"], source_path, payload.target_chars, payload.overlap_chars),
                "paragraph_merge": _paragraph_merge(loaded["text"], source_path, payload.target_chars, payload.max_chars),
                "heading_table_aware": _heading_table_aware(loaded["text"], source_path, payload.target_chars, payload.max_chars),
            }
            evaluated: list[dict[str, Any]] = []
            for name, chunks in strategies.items():
                metrics = _evaluate_chunks(chunks, payload.candidate_terms, payload.expected_terms, payload.noise_terms)
                evaluated.append({
                    "name": name,
                    "metrics": metrics,
                    "score": _score_metrics(metrics, payload.max_chars),
                    "chunks": chunks,
                    "sample_chunks": chunks[:3],
                })
            winner = max(evaluated, key=lambda item: float(item["score"]["score"]))
            documents.append({
                "document_id": document_id,
                "source_path": source_path,
                "status": "ok",
                "parser": loaded["parser"],
                "source_stats": loaded["stats"],
                "winner": winner["name"],
                "strategies": evaluated,
            })
        except Exception as exc:
            documents.append({
                "document_id": document_id,
                "source_path": document.source_path,
                "status": "error",
                "error": str(exc),
                "strategies": [],
            })

    evaluable = [document for document in documents if document.get("status") == "ok" and document.get("strategies")]
    candidate_deltas: list[float] = []
    noise_deltas: list[float] = []
    for document in evaluable:
        strategies = document["strategies"]
        baseline = next((item for item in strategies if item["name"] == "fixed_chars"), strategies[0])
        winner = next((item for item in strategies if item["name"] == document["winner"]), strategies[0])
        candidate_deltas.append(
            float(winner["metrics"]["candidate_coverage"]["rate"]) - float(baseline["metrics"]["candidate_coverage"]["rate"])
        )
        noise_deltas.append(int(winner["metrics"]["noise_hit_count"]) - int(baseline["metrics"]["noise_hit_count"]))

    avg_candidate_delta = round(sum(candidate_deltas) / max(len(candidate_deltas), 1), 6)
    avg_noise_delta = round(sum(noise_deltas) / max(len(noise_deltas), 1), 6)
    if not evaluable:
        recommendation = "inconclusive"
    elif avg_candidate_delta >= 0 and avg_noise_delta <= 0:
        recommendation = "recommended"
    elif avg_candidate_delta < 0 or avg_noise_delta > 5:
        recommendation = "not_recommended"
    else:
        recommendation = "inconclusive"

    report = {
        "run_id": run_id,
        "title": payload.title or "Adaptive Chunking UI Test",
        "created_at": datetime.now(UTC).isoformat(),
        "config": {
            **config,
            "artifact_dir": str(output_dir),
            "settings_artifact_dir": settings.ADAPTIVE_CHUNKING_ARTIFACT_DIR,
            "candidate_terms": payload.candidate_terms,
            "expected_terms": payload.expected_terms,
            "noise_terms": payload.noise_terms,
        },
        "summary": {
            "recommendation": recommendation,
            "document_count": len(documents),
            "evaluable_document_count": len(evaluable),
            "avg_candidate_coverage_delta": avg_candidate_delta,
            "avg_noise_delta": avg_noise_delta,
        },
        "documents": documents,
    }
    _write_json_report(report, output_dir)
    _write_markdown_report(report, output_dir)
    detail = get_adaptive_chunking_report(run_id)
    if detail is None:
        raise RuntimeError("Evaluation report was written but could not be read back")
    return detail
