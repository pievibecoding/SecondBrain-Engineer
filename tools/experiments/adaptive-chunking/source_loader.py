from __future__ import annotations

from pathlib import Path
from typing import Any


def _read_pdf(path: Path) -> str:
    try:
        import pdfplumber  # type: ignore
    except Exception as exc:  # pragma: no cover - depends on local env
        raise RuntimeError("pdfplumber is not installed") from exc

    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text(x_tolerance=1, y_tolerance=3) or "")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    try:
        from docx import Document  # type: ignore
    except Exception as exc:  # pragma: no cover - depends on local env
        raise RuntimeError("python-docx is not installed") from exc

    document = Document(path)
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"\n[Table {table_index}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def load_source_text(source_path: str, workspace: Path) -> dict[str, Any]:
    path = Path(source_path)
    if not path.is_absolute():
        path = workspace / path
    path = path.resolve()

    if not path.exists():
        raise FileNotFoundError(f"Source path not found: {path}")

    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        parser = "text"
    elif suffix == ".pdf":
        text = _read_pdf(path)
        parser = "pdfplumber"
    elif suffix == ".docx":
        text = _read_docx(path)
        parser = "python-docx"
    else:
        raise ValueError(f"Unsupported source extension: {suffix}")

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = text.splitlines()
    empty_lines = sum(1 for line in lines if not line.strip())
    broken_layout_terms = ["cirtceleotohP", "srosneS", "Bảng trang", "| --- |"]
    broken_layout_hits = sum(text.count(term) for term in broken_layout_terms)

    return {
        "source_path": str(path),
        "parser": parser,
        "text": text,
        "stats": {
            "char_count": len(text),
            "line_count": len(lines),
            "empty_line_ratio": round(empty_lines / max(len(lines), 1), 6),
            "broken_layout_hits": broken_layout_hits,
        },
    }

