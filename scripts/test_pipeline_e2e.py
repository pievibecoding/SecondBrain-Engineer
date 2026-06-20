"""
Test end-to-end: parse DOCX bằng python-docx → gửi vào LightRAG /documents/text
So sánh kết quả với lần ingest native trước đó.

Chạy: python scripts/test_pipeline_e2e.py
"""

import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import httpx

LIGHTRAG_URL = "http://localhost:9621"
DOCX_PATH = Path(r"D:\Robolinks Intern\Project\SecondBrain\local-nas\projects\Swarovski\Manual\Hướng dẫn vận hành_SATO.docx")
CORRELATION_ID = "test-pipeline-e2e-001"


def parse_docx(path: Path) -> tuple[str, int, int]:
    """Parse DOCX bằng python-docx DOM traversal."""
    import docx as _docx

    doc = _docx.Document(path)
    blocks = []
    table_index = 0
    para_index = 0
    table_count = 0

    for child in doc.element.body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag == "p":
            if para_index < len(doc.paragraphs):
                para = doc.paragraphs[para_index]
                para_index += 1
                text = para.text.strip()
                if text:
                    blocks.append(text)
        elif tag == "tbl":
            if table_index < len(doc.tables):
                table = doc.tables[table_index]
                table_index += 1
                rows = []
                for row in table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    rows.append("| " + " | ".join(cells) + " |")
                if rows:
                    table_count += 1
                    sep = "| " + " | ".join(["---"] * len(table.columns)) + " |"
                    rows.insert(1, sep)
                    blocks.append(f"\n[Bảng {table_count}]\n" + "\n".join(rows))

    plain = "\n\n".join(blocks)
    return plain, len(plain), table_count


def send_to_lightrag(text: str, file_source: str) -> dict:
    """Gửi text đã parse vào LightRAG /documents/text."""
    with httpx.Client(timeout=30) as client:
        resp = client.post(
            f"{LIGHTRAG_URL}/documents/text",
            json={
                "text": text,
                "file_source": file_source,
                "metadata": {
                    "source": "test_pipeline",
                    "parser": "python-docx",
                    "nas_path": str(DOCX_PATH),
                },
            },
            headers={"X-Correlation-ID": CORRELATION_ID},
        )
        resp.raise_for_status()
        return resp.json()


def poll_status(track_id: str, timeout: int = 300) -> dict:
    """Poll trạng thái ingest cho đến khi xong hoặc timeout."""
    deadline = time.time() + timeout
    while time.time() < deadline:
        with httpx.Client(timeout=10) as client:
            resp = client.get(f"{LIGHTRAG_URL}/documents/track_status/{track_id}")
            if resp.is_success:
                data = resp.json()
                status = data.get("status", "")
                print(f"  Status: {status}", end="\r")
                if status in {"processed", "failed", "error"}:
                    print()
                    return data
        time.sleep(3)
    return {"status": "timeout"}


def check_chunks(doc_id: str) -> list[dict]:
    """Query PostgreSQL để xem chunks đã tạo."""
    import subprocess
    result = subprocess.run(
        [
            "docker", "exec", "secondbrain-postgres-1",
            "psql", "-U", "secondbrain", "-d", "secondbrain",
            "-c", f"SELECT chunk_order_index, tokens, LEFT(content, 150) as preview FROM lightrag_doc_chunks WHERE full_doc_id = '{doc_id}' ORDER BY chunk_order_index;"
        ],
        capture_output=True, text=True, encoding="utf-8"
    )
    print(result.stdout)
    return []


def main():
    if not DOCX_PATH.exists():
        print(f"ERROR: File không tồn tại: {DOCX_PATH}")
        sys.exit(1)

    print("=" * 70)
    print("TEST: python-docx → LightRAG /documents/text")
    print("=" * 70)

    # Step 1: Parse
    print("\n[1/3] Parsing DOCX với python-docx...")
    t0 = time.time()
    text, char_count, table_count = parse_docx(DOCX_PATH)
    parse_time = time.time() - t0
    print(f"  ✅ Parse xong: {char_count:,} chars, {table_count} tables, {parse_time:.1f}s")
    print(f"  Preview: {text[:200].replace(chr(10), ' ')}")

    # Step 2: Gửi vào LightRAG
    print("\n[2/3] Gửi text vào LightRAG /documents/text...")
    try:
        t1 = time.time()
        response = send_to_lightrag(text, file_source="SATO_docx_python-docx_test")
        send_time = time.time() - t1
        print(f"  ✅ Accepted: {response}")
        track_id = response.get("track_id") or response.get("id", "")
    except httpx.HTTPStatusError as e:
        print(f"  ❌ HTTP Error: {e.response.status_code} — {e.response.text}")
        sys.exit(1)
    except Exception as e:
        print(f"  ❌ Error: {e}")
        sys.exit(1)

    # Step 3: Poll status
    if track_id:
        print(f"\n[3/3] Chờ LightRAG xử lý (track_id: {track_id})...")
        t2 = time.time()
        final = poll_status(track_id, timeout=300)
        total_time = time.time() - t2
        status = final.get("status", "unknown")

        if status == "processed":
            print(f"  ✅ Ingest thành công! Tổng thời gian: {total_time:.0f}s")
            doc_id = final.get("id") or track_id
            print(f"  Doc ID: {doc_id}")
            print(f"\n📊 CHUNKS:")
            check_chunks(doc_id)
        elif status == "failed":
            print(f"  ❌ Ingest thất bại: {final.get('error_msg', 'unknown error')}")
        else:
            print(f"  ⚠️  Status: {status} — {final}")
    else:
        print("  ⚠️  Không có track_id, không thể poll status")

    print("\n" + "=" * 70)
    print("So sánh với Test 1 (LightRAG native):")
    print("  Native:     14 chunks, ~9,680 tokens, thiếu ~50% text")
    print(f"  python-docx: ? chunks, {char_count:,} chars, {table_count} tables đầy đủ")
    print("=" * 70)


if __name__ == "__main__":
    main()
